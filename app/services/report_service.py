"""
Report Service — ReportLab PDF Slips & OpenPyXL Excel Export.
Demonstrates TQM Principle: Documentation & Quality Reporting.
"""

import os
from datetime import datetime
from pathlib import Path
from app.config import REPORTS_DIR
from app.core.logger import get_logger

logger = get_logger("reports")


def _ensure_reports_dir():
    """Ensures the reports directory exists."""
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)


def generate_payslip_pdf(staff_name: str, period: str, gross: float,
                         deductions: float, net: float, regular_hours: float,
                         overtime_hours: float, hourly_rate: float) -> str:
    """
    Generates a professional PDF payslip using ReportLab.
    Returns the file path of the generated PDF.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    _ensure_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"payslip_{staff_name.replace(' ', '_')}_{period.replace('/', '-')}_{timestamp}.pdf"
    filepath = str(REPORTS_DIR / filename)

    doc = SimpleDocTemplate(filepath, pagesize=A4,
                            topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    elements.append(Paragraph("CampusLib Pro - Staff Pay Slip", styles["Title"]))
    elements.append(Spacer(1, 10 * mm))

    # Pay Details Table
    data = [
        ["Staff Name", staff_name],
        ["Pay Period", period],
        ["Regular Hours", f"{regular_hours:.1f} hrs"],
        ["Overtime Hours", f"{overtime_hours:.1f} hrs"],
        ["Hourly Rate", f"Rs. {hourly_rate:.2f}"],
        ["", ""],
        ["Gross Salary", f"Rs. {gross:.2f}"],
        ["Deductions (5%)", f"Rs. {deductions:.2f}"],
        ["Net Salary", f"Rs. {net:.2f}"],
    ]

    table = Table(data, colWidths=[60 * mm, 80 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DBEAFE")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 15 * mm))

    # Footer
    elements.append(Paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | TQM Automated Payroll System",
        styles["Italic"]
    ))

    doc.build(elements)
    logger.info(f"Pay slip PDF generated: {filepath}")
    return filepath


def generate_catalog_excel(books_data: list) -> str:
    """
    Exports book catalog to an Excel file using OpenPyXL.
    Returns the file path of the generated Excel file.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    _ensure_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"catalog_report_{timestamp}.xlsx"
    filepath = str(REPORTS_DIR / filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Book Catalog"

    # Header styling
    header_font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = ["ID", "Title", "Author", "Category", "ISBN", "Year",
               "Total Copies", "Available", "Defect Flag"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Data rows
    for row_idx, book in enumerate(books_data, 2):
        values = [
            book.get("id", ""), book.get("title", ""), book.get("author", ""),
            book.get("category", ""), book.get("isbn", ""), book.get("year", ""),
            book.get("total_copies", ""), book.get("available_copies", ""),
            "Yes" if book.get("defect_flag", False) else "No",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col not in [2, 3] else "left")

    # Auto-fit column widths
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    wb.save(filepath)
    logger.info(f"Catalog Excel report generated: {filepath}")
    return filepath


def generate_quality_audit_excel(quality_logs: list) -> str:
    """
    Exports TQM quality log entries to an Excel audit trail.
    Returns the file path of the generated Excel file.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    _ensure_reports_dir()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"quality_audit_{timestamp}.xlsx"
    filepath = str(REPORTS_DIR / filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "TQM Quality Audit"

    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = ["ID", "Timestamp", "Module", "Action", "Status", "Latency (ms)", "Details", "Remedial Action"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row_idx, log in enumerate(quality_logs, 2):
        values = [
            log.get("id", ""), str(log.get("timestamp", "")), log.get("module", ""),
            log.get("action", ""), log.get("status", ""),
            log.get("execution_time_ms", 0), log.get("details", ""),
            log.get("remedial_action", ""),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 45)

    wb.save(filepath)
    logger.info(f"Quality audit Excel report generated: {filepath}")
    return filepath
